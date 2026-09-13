"""Report generator for Rehab RCH Agent v2.

Builds branded .docx documents (weekly/monthly/annual reports,
announcements, meeting minutes) and saves them locally under
``output/`` mirroring the Google Drive folder structure::

    output/Reports/Weekly/Weekly Summary - 13 Sep 2026.docx

Drive upload is handled by :mod:`google_drive`.
"""

from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path
from typing import Optional

from config import settings

FOOTER_NOTE = "Operational use only — no patient information. Draft for review — Rehabilitation Department, RCH."


def slugify(text: str, max_len: int = 60) -> str:
    text = re.sub(r"[^\w\s-]", "", text).strip()
    text = re.sub(r"[\s_-]+", " ", text).strip()
    return text[:max_len] or "Report"


def dated_filename(prefix: str, when: Optional[datetime] = None) -> str:
    when = when or datetime.now()
    return f"{prefix} - {when.strftime('%d %b %Y')}.docx"


def local_output_path(*parts: str, filename: str) -> Path:
    path = settings.output_dir.joinpath(*parts)
    path.mkdir(parents=True, exist_ok=True)
    return path / filename


def _style_document(doc, title: str, subtitle: str, prepared_by: str, when: datetime):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    # Header block
    h0 = doc.add_paragraph()
    h0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = h0.add_run(settings.department_name)
    run.bold = True
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor(0x1A, 0x5C, 0x38)

    h1 = doc.add_paragraph()
    h1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = h1.add_run(title)
    run.bold = True
    run.font.size = Pt(16)

    if subtitle:
        sub = doc.add_paragraph()
        sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = sub.add_run(subtitle)
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = meta.add_run(
        f"Date: {when.strftime('%d %B %Y')}    |    Prepared by: {prepared_by}"
    )
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    doc.add_paragraph("─" * 60)

    # Footer
    section = doc.sections[0]
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run(FOOTER_NOTE)
    run.font.size = Pt(7)
    run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)


def _add_markdownish_body(doc, body: str) -> None:
    """Render Gemini markdown-ish output into docx paragraphs.

    Supports: '#'-headings, '**bold**' lines, '- '/ '• ' bullets,
    numbered lists, and plain paragraphs.
    """
    from docx.shared import Pt

    for raw_line in body.splitlines():
        line = raw_line.strip()
        if not line:
            doc.add_paragraph("")
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:].strip("* ").strip(), level=3)
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip("* ").strip(), level=2)
        elif line.startswith("# "):
            doc.add_heading(line[2:].strip("* ").strip(), level=1)
        elif re.match(r"^[-•]\s+", line):
            doc.add_paragraph(re.sub(r"^[-•]\s+", "", line).strip("* ").strip(), style="List Bullet")
        elif re.match(r"^\d+[.)]\s+", line):
            doc.add_paragraph(re.sub(r"^\d+[.)]\s+", "", line).strip("* ").strip(), style="List Number")
        elif line.startswith("**") and line.endswith("**") and len(line) > 4:
            p = doc.add_paragraph()
            run = p.add_run(line.strip("* ").strip())
            run.bold = True
            run.font.size = Pt(12)
        else:
            # inline bold cleanup: **x** -> x (keep simple, avoid runs complexity)
            clean = re.sub(r"\*\*(.+?)\*\*", r"\1", line)
            clean = re.sub(r"^Draft for review\s*$", "Draft for review", clean)
            doc.add_paragraph(clean)


def build_docx(
    title: str,
    body: str,
    prepared_by: str = "Rehab RCH Agent",
    subtitle: str = "",
    when: Optional[datetime] = None,
) -> "Path | object":
    """Build a branded docx Document and return (path not saved).

    Use the typed helpers below (which save to disk) in normal flows.
    """
    import docx

    when = when or datetime.now()
    doc = docx.Document()
    _style_document(doc, title, subtitle, prepared_by, when)
    _add_markdownish_body(doc, body)
    return doc


def _save(doc, path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path


# -- typed builders ---------------------------------------------------------

def build_weekly_report(body: str, prepared_by: str = "Rehab RCH Agent",
                        when: Optional[datetime] = None) -> Path:
    when = when or datetime.now()
    doc = build_docx(
        "Weekly Rehabilitation Operations Summary",
        body, prepared_by,
        subtitle=f"Week of {when.strftime('%d %B %Y')}", when=when,
    )
    return _save(doc, local_output_path("Reports", "Weekly", filename=dated_filename("Weekly Summary", when)))


def build_monthly_report(body: str, prepared_by: str = "Rehab RCH Agent",
                         when: Optional[datetime] = None) -> Path:
    when = when or datetime.now()
    doc = build_docx(
        "Monthly Rehabilitation Operations Report",
        body, prepared_by,
        subtitle=when.strftime("%B %Y"), when=when,
    )
    return _save(doc, local_output_path("Reports", "Monthly", filename=dated_filename("Monthly Report", when)))


def build_annual_report(body: str, prepared_by: str = "Rehab RCH Agent",
                        when: Optional[datetime] = None) -> Path:
    when = when or datetime.now()
    doc = build_docx(
        "Annual Rehabilitation Operations Report",
        body, prepared_by,
        subtitle=when.strftime("%Y"), when=when,
    )
    return _save(doc, local_output_path("Reports", "Annual", filename=dated_filename("Annual Report", when)))


def build_announcement(body: str, subject: str = "Department Announcement",
                       prepared_by: str = "Rehab RCH Agent",
                       when: Optional[datetime] = None) -> Path:
    when = when or datetime.now()
    doc = build_docx(f"Announcement — {slugify(subject, 50)}", body, prepared_by, when=when)
    return _save(
        doc,
        local_output_path("Announcements", filename=dated_filename(f"Announcement - {slugify(subject, 40)}", when)),
    )


def build_meeting_minutes(body: str, title: str = "Staff Meeting",
                          prepared_by: str = "Rehab RCH Agent",
                          when: Optional[datetime] = None) -> Path:
    when = when or datetime.now()
    doc = build_docx(f"Meeting Minutes — {slugify(title, 50)}", body, prepared_by, when=when)
    return _save(
        doc,
        local_output_path("Meeting Minutes", filename=dated_filename(f"Meeting Minutes - {slugify(title, 40)}", when)),
    )


def build_operational_report(body: str, title: str = "Operational Report",
                             prepared_by: str = "Rehab RCH Agent",
                             when: Optional[datetime] = None) -> Path:
    when = when or datetime.now()
    doc = build_docx(title, body, prepared_by, when=when)
    return _save(
        doc,
        local_output_path("Reports", filename=dated_filename(slugify(title, 40), when)),
    )
